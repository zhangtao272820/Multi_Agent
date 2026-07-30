"""Workspace Office 文档读写（docx / xlsx），仅限 WORKSPACE_DIR。"""
from __future__ import annotations

from pathlib import Path

from app.tools.common import _tool_err, _tool_ok
from app.tools.files import _rel_display, _resolve_workspace_path


def _ext(path: Path) -> str:
    return path.suffix.lower().lstrip(".")


def read_office_document(file_path: str, sheet: str = "") -> str:
    """读取工作区内 docx/xlsx 为可读文本。"""
    safe = _resolve_workspace_path(file_path, allow_missing=False)
    if safe is None or not safe.is_file():
        return _tool_err("文件不存在或路径非法。", code="path_escape")
    kind = _ext(safe)
    try:
        if kind == "docx":
            from docx import Document

            doc = Document(str(safe))
            paras = [p.text for p in doc.paragraphs if (p.text or "").strip()]
            tables_txt: list[str] = []
            for ti, table in enumerate(doc.tables):
                rows = []
                for row in table.rows:
                    rows.append("\t".join((c.text or "").strip() for c in row.cells))
                tables_txt.append(f"[table {ti}]\n" + "\n".join(rows))
            body = "\n".join(paras)
            if tables_txt:
                body = (body + "\n\n" if body else "") + "\n\n".join(tables_txt)
            if len(body) > 32_000:
                body = body[:32_000] + f"\n\n…[已截断，原文约更长]"
            return _tool_ok(
                body or "(空文档)",
                data={"file_path": _rel_display(safe), "format": "docx", "chars": len(body)},
            )
        if kind in ("xlsx", "xlsm"):
            from openpyxl import load_workbook

            wb = load_workbook(str(safe), read_only=True, data_only=True)
            names = list(wb.sheetnames)
            target = sheet.strip() if sheet else (names[0] if names else "")
            if not target or target not in wb.sheetnames:
                wb.close()
                return _tool_err(
                    f"工作表不存在。可用: {', '.join(names)}",
                    data={"sheets": names},
                    code="sheet_not_found",
                )
            ws = wb[target]
            lines: list[str] = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= 200:
                    lines.append("…[已截断，最多 200 行]")
                    break
                cells = ["" if c is None else str(c) for c in row]
                if any(cells):
                    lines.append("\t".join(cells))
            wb.close()
            body = "\n".join(lines) or "(空表)"
            return _tool_ok(
                body,
                data={
                    "file_path": _rel_display(safe),
                    "format": "xlsx",
                    "sheet": target,
                    "sheets": names,
                    "rows": len(lines),
                },
            )
        return _tool_err(
            f"不支持的格式 .{kind}，仅支持 docx/xlsx。",
            data={"file_path": _rel_display(safe)},
            code="unsupported_format",
        )
    except ImportError as e:
        return _tool_err(
            f"缺少 Office 依赖（python-docx / openpyxl）: {e}",
            code="missing_dependency",
        )
    except Exception as e:
        return _tool_err(f"读取 Office 文档失败: {e}", code="read_office_failed")


def write_office_document(
    file_path: str,
    content: str = "",
    format: str = "",
    rows: list | None = None,
    sheet: str = "Sheet1",
) -> str:
    """写入工作区 docx（段落文本）或 xlsx（二维 rows）。"""
    safe = _resolve_workspace_path(file_path, allow_missing=True)
    if safe is None or safe == _resolve_workspace_path(""):
        return _tool_err("非法文件路径。", code="path_escape")
    kind = (_ext(safe) or str(format or "").lower().strip().lstrip("."))
    if kind not in ("docx", "xlsx"):
        return _tool_err("format/后缀须为 docx 或 xlsx。", code="unsupported_format")
    try:
        safe.parent.mkdir(parents=True, exist_ok=True)
        if kind == "docx":
            from docx import Document

            doc = Document()
            text = content if content is not None else ""
            for para in str(text).split("\n"):
                doc.add_paragraph(para)
            if not safe.suffix:
                safe = safe.with_suffix(".docx")
            doc.save(str(safe))
            return _tool_ok(
                f"已写入 Word 文档: {_rel_display(safe)}",
                data={"file_path": _rel_display(safe), "format": "docx"},
            )
        # xlsx
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = (sheet or "Sheet1")[:31] or "Sheet1"
        matrix = rows if isinstance(rows, list) and rows else None
        if matrix is None:
            # content 按行、制表符分列
            matrix = []
            for line in str(content or "").splitlines():
                matrix.append(line.split("\t") if "\t" in line else [line])
        for r in matrix:
            if isinstance(r, (list, tuple)):
                ws.append(list(r))
            else:
                ws.append([r])
        if not safe.suffix:
            safe = safe.with_suffix(".xlsx")
        wb.save(str(safe))
        return _tool_ok(
            f"已写入 Excel 文档: {_rel_display(safe)}",
            data={
                "file_path": _rel_display(safe),
                "format": "xlsx",
                "sheet": ws.title,
                "row_count": len(matrix),
            },
        )
    except ImportError as e:
        return _tool_err(
            f"缺少 Office 依赖（python-docx / openpyxl）: {e}",
            code="missing_dependency",
        )
    except Exception as e:
        return _tool_err(f"写入 Office 文档失败: {e}", code="write_office_failed")
